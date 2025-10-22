#!/usr/bin/env python3
"""
Gebruik:
$ python generate_quiz.py vragen.xlsx
Leest een Excel/CSV met kolommen:
leeftijd, categorie, kleur, type, vraag, antwoord, opties, plaatje, hint, uitleg, moeilijk
Schrijft per leeftijd/categorie een JSON-bestand naar ../public/data/groep<leeftijd>_<categorie>.json
"""

import sys, json, csv, pathlib, re
from openpyxl import load_workbook

def slugify(s):
    return re.sub(r'[^a-z0-9]+','_',s.lower()).strip('_')

def read_rows(path):
    if path.suffix.lower()=='.csv':
        with open(path, newline='', encoding='utf-8') as f:
            yield from csv.DictReader(f)
    else:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c for c in next(ws.iter_rows(values_only=True))]
        for row in ws.iter_rows(values_only=True):
            yield dict(zip(headers, row))

def main():
    if len(sys.argv)<2:
        sys.exit('geef invoerbestand')
    inp = pathlib.Path(sys.argv[1])
    outdir = inp.parent.parent / 'public' / 'data'
    outdir.mkdir(parents=True, exist_ok=True)
    vragen = {}
    for r in read_rows(inp):
        leeftijd = int(r['leeftijd'])
        categorie = str(r['categorie']).strip()
        cat_slug = slugify(categorie)
        key = (leeftijd, cat_slug)
        if key not in vragen:
            vragen[key] = []
        typ = r.get('type','mcq')
        item = dict(
            id       = f"g{leeftijd}_{cat_slug}_{len(vragen[key])+1:03}",
            leeftijd = leeftijd,
            categorie= categorie,
            kleur    = r.get('kleur') or None,
            type     = typ,
            vraag    = r['vraag'],
            antwoord = r['antwoord'],
            moeilijk = int(r.get('moeilijk',1)),
            plaatje  = r.get('plaatje') or None,
            hint     = r.get('hint') or None,
            uitleg   = r.get('uitleg') or None
        )
        if typ=='mcq' and r.get('opties'):
            item['opties']=[o.strip() for o in str(r['opties']).split(',')]
        vragen[key].append(item)
    # Schrijf per categorie/leeftijd een bestand
    for (leeftijd,cat_slug), items in vragen.items():
        out = outdir / f"groep{leeftijd}_{cat_slug}.json"
        with open(out,'w',encoding='utf-8') as f:
            json.dump(items, f, ensure_ascii=False, indent=2)
        print(f"Geschreven: {out}")

if __name__=='__main__':
    main() 